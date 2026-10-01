#!/usr/bin/env python3
"""
Check which learners appear in multiple sheets and identify the issue.
"""

from openpyxl import load_workbook

SPREADSHEET_FILE = "REMEDIAL_PAYMENT_BALANCES_2026.3_FIXED(1).xlsx"

def main():
    print("=" * 80)
    print("CHECKING FOR DUPLICATE LEARNERS ACROSS SHEETS")
    print("=" * 80)
    print()
    
    workbook = load_workbook(SPREADSHEET_FILE)
    
    # Track learners by admission number
    learner_locations = {}
    
    for sheet_name in workbook.sheetnames:
        worksheet = workbook[sheet_name]
        print(f"\nProcessing sheet: {sheet_name}")
        
        for row_idx in range(5, worksheet.max_row + 1):
            # Check column B for admission numbers (standard structure)
            adm_cell_b = worksheet.cell(row=row_idx, column=2).value
            # Check column D for admission numbers (Form 4 structure)
            adm_cell_d = worksheet.cell(row=row_idx, column=4).value
            
            adm_value = adm_cell_b if adm_cell_b else adm_cell_d
            
            if adm_value and str(adm_value).strip().isdigit():
                adm_num = str(adm_value).strip()
                name_cell = worksheet.cell(row=row_idx, column=3 if adm_cell_b else 5).value
                
                if name_cell and str(name_cell).strip() and str(name_cell).strip() != 'NAME':
                    name = str(name_cell).strip()
                    
                    if adm_num not in learner_locations:
                        learner_locations[adm_num] = []
                    
                    learner_locations[adm_num].append({
                        'sheet': sheet_name,
                        'row': row_idx,
                        'name': name
                    })
    
    # Find duplicates
    duplicates = {adm: locs for adm, locs in learner_locations.items() if len(locs) > 1}
    
    print("\n" + "=" * 80)
    print(f"FOUND {len(duplicates)} LEARNERS IN MULTIPLE SHEETS")
    print("=" * 80)
    
    # Focus on Form 2 2025 and Form 3 2026 duplicates
    form2_form3_duplicates = []
    
    for adm_num, locations in duplicates.items():
        sheets = [loc['sheet'] for loc in locations]
        if 'FORM 2 2025' in sheets and 'FORM 3 2026' in sheets:
            form2_form3_duplicates.append((adm_num, locations))
    
    print(f"\nLearners in BOTH 'FORM 2 2025' AND 'FORM 3 2026': {len(form2_form3_duplicates)}")
    print()
    
    if form2_form3_duplicates:
        print("Sample learners (first 10):")
        for adm_num, locations in form2_form3_duplicates[:10]:
            print(f"\n  Admission {adm_num}:")
            for loc in locations:
                print(f"    - {loc['sheet']} (Row {loc['row']}): {loc['name']}")
    
    # Check a specific learner's data in both sheets
    print("\n" + "=" * 80)
    print("DETAILED CHECK: Comparing data for a sample learner")
    print("=" * 80)
    
    if form2_form3_duplicates:
        sample_adm = form2_form3_duplicates[0][0]
        print(f"\nChecking admission {sample_adm}:")
        
        for sheet_name in ['FORM 2 2025', 'FORM 3 2026']:
            if sheet_name in workbook.sheetnames:
                worksheet = workbook[sheet_name]
                print(f"\n  Sheet: {sheet_name}")
                
                # Find the learner
                for row_idx in range(5, worksheet.max_row + 1):
                    adm_cell = worksheet.cell(row=row_idx, column=2).value
                    if adm_cell and str(adm_cell).strip() == sample_adm:
                        name = worksheet.cell(row=row_idx, column=3).value
                        total_owed = worksheet.cell(row=row_idx, column=5).value
                        
                        print(f"    Row: {row_idx}")
                        print(f"    Name: {name}")
                        print(f"    Total Owed: {total_owed}")
                        
                        # Check for payments
                        payments = []
                        for col_idx in range(8, 20, 2):  # Check first few payment columns
                            amount = worksheet.cell(row=row_idx, column=col_idx).value
                            date = worksheet.cell(row=row_idx, column=col_idx + 1).value
                            if amount:
                                payments.append((amount, date))
                        
                        print(f"    Payments found: {len(payments)}")
                        if payments:
                            for i, (amt, dt) in enumerate(payments[:3], 1):
                                print(f"      Payment {i}: {amt} on {dt}")
                        
                        break
    
    workbook.close()
    print()


if __name__ == "__main__":
    main()
