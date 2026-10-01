"""
Attendance Tracker - Records teacher attendance
© 2026 Kirembe Secondary School
"""
import openpyxl
from datetime import datetime
from typing import Dict, List, Optional


class AttendanceTracker:
    """Tracks and records teacher attendance"""
    
    def __init__(self, excel_file: str = "TEACHER_PAYROLL.xlsx"):
        self.excel_file = excel_file
        self.weekday_rate = 200  # KES per weekday
        self.weekend_rate = 500  # KES per weekend
    
    def record_attendance(self, teacher_id: int, term: int, week: int, 
                         days_attended: List[str] = None, days: Dict[str, bool] = None,
                         teacher_name: str = None) -> Dict:
        """
        Record attendance for a teacher.
        
        Args:
            teacher_id: Teacher's ID number
            term: Term number (1, 2, or 3)
            week: Week number (1-12)
            days_attended: List of days attended (e.g., ['monday', 'tuesday', 'saturday'])
                          - This is what the web app sends
            days: Dictionary of days attended (alternative format):
                  {'monday': True, 'tuesday': False, ..., 'saturday': True}
            teacher_name: Teacher's name (optional - will be fetched if not provided)
        
        Returns:
            Dictionary with success status and calculated pay
        """
        try:
            # Handle both parameter formats
            if days_attended is None and days is None:
                days_attended = []
            
            # Convert list format to dictionary if needed
            if days_attended is not None:
                # Normalize day names to lowercase
                days_lower = [d.lower().strip() for d in days_attended]
                days_dict = {
                    'monday': 'monday' in days_lower,
                    'tuesday': 'tuesday' in days_lower,
                    'wednesday': 'wednesday' in days_lower,
                    'thursday': 'thursday' in days_lower,
                    'friday': 'friday' in days_lower,
                    'saturday': 'saturday' in days_lower,
                    'saturdaytp': 'saturdaytp' in days_lower
                }
            else:
                # Use the dictionary format provided
                days_dict = days if days else {}
            
            # Get teacher name if not provided
            if teacher_name is None:
                teacher_name = self._get_teacher_name(teacher_id)
            
            # Calculate attendance
            weekday_count = sum([
                days_dict.get('monday', False),
                days_dict.get('tuesday', False),
                days_dict.get('wednesday', False),
                days_dict.get('thursday', False),
                days_dict.get('friday', False)
            ])
            
            weekend_count = 1 if days_dict.get('saturday', False) else 0
            
            weekday_pay = weekday_count * self.weekday_rate
            weekend_pay = weekend_count * self.weekend_rate
            
            # Second weekend lesson: SaturdayTP at 300 KES
            if days_dict.get('saturdaytp', False):
                weekend_pay += 300
            
            total_pay = weekday_pay + weekend_pay
            
            # Build days string for storage
            days_list = []
            for day in ['monday', 'tuesday', 'wednesday', 'thursday', 'friday', 'saturday', 'saturdaytp']:
                if days_dict.get(day, False):
                    label = 'SaturdayTP' if day == 'saturdaytp' else day.capitalize()
                    days_list.append(label)
            days_string = ', '.join(days_list)
            
            # Open Excel file
            wb = openpyxl.load_workbook(self.excel_file)
            ws = wb['Attendance']
            
            # Find next empty row
            next_row = ws.max_row + 1
            
            # Generate record ID
            record_id = f"T{term}W{week}T{teacher_id}"
            
            # Write attendance record
            ws[f'A{next_row}'] = record_id
            ws[f'B{next_row}'] = term
            ws[f'C{next_row}'] = week
            ws[f'D{next_row}'] = teacher_id
            ws[f'E{next_row}'] = teacher_name
            ws[f'F{next_row}'] = days_string  # Days attended as string
            ws[f'G{next_row}'] = weekday_pay
            ws[f'H{next_row}'] = weekend_pay
            ws[f'I{next_row}'] = total_pay
            ws[f'J{next_row}'] = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
            
            # Save workbook
            wb.save(self.excel_file)
            wb.close()
            
            return {
                'success': True,
                'message': f'Attendance recorded for {teacher_name}',
                'record_id': record_id,
                'teacher_id': teacher_id,
                'teacher_name': teacher_name,
                'term': term,
                'week': week,
                'days_attended': days_string,
                'weekday_count': weekday_count,
                'weekend_count': weekend_count,
                'weekday_pay': weekday_pay,
                'weekend_pay': weekend_pay,
                'total_pay': total_pay
            }
            
        except Exception as e:
            return {
                'success': False,
                'message': f'Error recording attendance: {str(e)}'
            }
    
    def _get_teacher_name(self, teacher_id: int) -> str:
        """Get teacher name from Teachers sheet"""
        try:
            wb = openpyxl.load_workbook(self.excel_file)
            ws = wb['Teachers']
            
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0] == teacher_id or str(row[0]) == str(teacher_id):
                    wb.close()
                    return row[1] if row[1] else f"Teacher {teacher_id}"
            
            wb.close()
        except Exception:
            pass
        return f"Teacher {teacher_id}"
    
    def get_all_attendance(self, term: int = None, week: int = None) -> List[Dict]:
        """
        Get all attendance records - ALIAS for compatibility.
        
        This method is called by the web app.
        """
        return self.get_all_attendance_records(term=term, week=week)
    
    def get_all_attendance_records(self, term: int = None, week: int = None) -> List[Dict]:
        """
        Get attendance records, optionally filtered by term and week.
        
        Args:
            term: Optional term number to filter by
            week: Optional week number to filter by
        
        Returns:
            List of attendance records
        """
        try:
            wb = openpyxl.load_workbook(self.excel_file)
            ws = wb['Attendance']
            
            records = []
            # Start from row 2 (skip header)
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0] is None:  # Skip empty rows (don't stop — deleted records leave gaps)
                    continue
                
                record = {
                    'record_id': row[0],
                    'term': row[1],
                    'week': row[2],
                    'teacher_id': row[3],
                    'teacher_name': row[4],
                    'days_attended': row[5],
                    'weekday_pay': row[6],
                    'weekend_pay': row[7],
                    'total_pay': row[8],
                    'date_recorded': row[9]
                }
                
                # Apply filters
                if term is not None and record['term'] != term:
                    continue
                if week is not None and record['week'] != week:
                    continue
                
                records.append(record)
            
            wb.close()
            return records
            
        except FileNotFoundError:
            return []
        except Exception as e:
            print(f"Error reading attendance: {e}")
            return []
    
    def delete_attendance_record(self, record_id: str) -> Dict:
        """
        Delete an attendance record by its ID.
        
        Args:
            record_id: The record ID to delete (format: T{term}W{week}T{teacher_id})
        
        Returns:
            Dictionary with success status and message
        """
        try:
            wb = openpyxl.load_workbook(self.excel_file)
            ws = wb['Attendance']
            
            # Find the record
            record_found = False
            row_to_delete = None
            
            for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=False), start=2):
                if row[0].value == record_id:
                    record_found = True
                    row_to_delete = row_idx
                    break
            
            if not record_found:
                wb.close()
                return {
                    'success': False,
                    'message': f'Record {record_id} not found'
                }
            
            # Delete the row
            ws.delete_rows(row_to_delete, 1)
            
            # Save workbook
            wb.save(self.excel_file)
            wb.close()
            
            return {
                'success': True,
                'message': f'Attendance record {record_id} deleted successfully'
            }
            
        except Exception as e:
            return {
                'success': False,
                'message': f'Error deleting attendance record: {str(e)}'
            }
    
    def get_total_payroll(self, term: int = None) -> Dict:
        """
        Get total payroll summary.
        
        Args:
            term: Optional term filter
            
        Returns:
            Dictionary with payroll totals
        """
        records = self.get_all_attendance_records(term=term)
        
        total_weekday_pay = sum(r.get('weekday_pay', 0) or 0 for r in records)
        total_weekend_pay = sum(r.get('weekend_pay', 0) or 0 for r in records)
        total_pay = sum(r.get('total_pay', 0) or 0 for r in records)
        
        return {
            'total_records': len(records),
            'total_weekday_pay': total_weekday_pay,
            'total_weekend_pay': total_weekend_pay,
            'total_pay': total_pay,
            'term': term
        }