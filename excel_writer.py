"""ExcelWriter for persisting domain objects to Excel spreadsheet."""

import logging
from datetime import datetime
from decimal import Decimal
from typing import List
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from src.domain.learner import Learner
from src.domain.payment_record import PaymentRecord
from src.utils.logger import setup_logger

logger = setup_logger(__name__)


class ExcelWriter:
    """
    Handles writing domain objects back to Excel spreadsheet format.
    
    The actual Excel file structure:
    - Column A: NO. (row number)
    - Column B: ADM (admission number)
    - Column C: NAME
    - Column D: ARREAS (not used)
    - Column E: DEBT (total owed)
    - Columns F-G: Not used
    - Columns H+: Payment records (monthly payments)
    """
    
    def write_learner(self, learner: Learner, file_path: str) -> dict:
        """
        Update a learner's row in the spreadsheet.
        
        Finds the learner by admission number and updates their entire row
        with current data. Amounts only are written to the remedial balances file.
        Payment dates are logged separately to PAYMENT_DATES_LOG.xlsx.
        
        Args:
            learner: The Learner object to write
            file_path: Path to the Excel file
            
        Returns:
            dict with 'success' (bool) and 'message' (str) keys
        """
        try:
            logger.info(f"Attempting to write learner {learner.admission_number} to {file_path}")
            workbook = load_workbook(file_path)
            
            # Use the learner's source sheet if available, otherwise use active sheet
            stream_suffix = None
            if hasattr(learner, 'source_sheet') and learner.source_sheet:
                # Extract base sheet name (remove stream suffix like " - Achievers" or " - Champions")
                base_sheet_name = learner.source_sheet
                if ' - Achievers' in base_sheet_name:
                    stream_suffix = '- Achievers'
                    base_sheet_name = base_sheet_name.replace(' - Achievers', '')
                elif ' - Champions' in base_sheet_name:
                    stream_suffix = '- Champions'
                    base_sheet_name = base_sheet_name.replace(' - Champions', '')
                
                if base_sheet_name in workbook.sheetnames:
                    worksheet = workbook[base_sheet_name]
                    logger.info(f"Using source sheet: {base_sheet_name} (from {learner.source_sheet})")
                else:
                    logger.warning(f"Source sheet '{base_sheet_name}' not found, using active sheet")
                    worksheet = workbook.active
            else:
                worksheet = workbook.active
            
            # Find the row for this learner (pass stream info to narrow search)
            row_index = self._find_learner_row(worksheet, learner.admission_number, worksheet.title, stream_suffix)
            
            if row_index is None:
                error_msg = (
                    f"Learner with admission number '{learner.admission_number}' "
                    f"not found in spreadsheet. The learner may have been removed "
                    f"or the admission number may have changed."
                )
                logger.error(error_msg)
                return {
                    'success': False,
                    'message': error_msg
                }
            
            # Prepare the data for the row - pass full payment records with timestamps
            # so update_row can write each payment to the correct month column
            row_data = {
                'name': learner.name,
                'admission_number': learner.admission_number,
                'total_owed': float(learner.total_owed),
                'payments': self.format_payment_records(learner.payment_records),
                'payment_records': learner.payment_records  # Full records with timestamps
            }
            
            # Update the row
            result = self.update_row(row_index, row_data, worksheet)
            
            if result['success']:
                workbook.save(file_path)
                workbook.close()
                logger.info(f"Successfully saved learner {learner.admission_number} to spreadsheet")
                
                # Log payment dates to PAYMENT_DATES_LOG.xlsx
                import os
                log_file = os.path.join(os.path.dirname(file_path), 'PAYMENT_DATES_LOG.xlsx')
                if not os.path.isabs(log_file):
                    log_file = 'PAYMENT_DATES_LOG.xlsx'
                self._log_payment_dates(learner, log_file)
            else:
                logger.error(f"Failed to update row for learner {learner.admission_number}: {result['message']}")
            
            return result
            
        except FileNotFoundError:
            error_msg = (
                f"Cannot find payment database file: {file_path}. "
                "The file may have been moved or deleted. "
                "Please ensure the file exists in the correct location."
            )
            logger.error(error_msg)
            return {
                'success': False,
                'message': error_msg
            }
        except PermissionError:
            error_msg = (
                f"Cannot access payment database file: Permission denied. "
                f"The file '{file_path}' may be open in Excel or another program. "
                "Please close the file and try again. Changes have been saved in memory only."
            )
            logger.error(error_msg)
            return {
                'success': False,
                'message': error_msg
            }
        except InvalidFileException:
            error_msg = (
                f"Payment database file has invalid format. "
                "The file may be corrupted or not a valid Excel (.xlsx) file. "
                "Please check the file and try again."
            )
            logger.error(error_msg)
            return {
                'success': False,
                'message': error_msg
            }
        except Exception as e:
            error_msg = (
                f"Cannot write to payment database file: {str(e)}. "
                "An unexpected error occurred. Changes have been saved in memory only."
            )
            logger.error(f"{error_msg} - Full error: {repr(e)}")
            return {
                'success': False,
                'message': error_msg
            }
    
    def format_payment_records(self, payments: List[PaymentRecord]) -> List:
        """
        Convert payment records to Excel format.
        
        Formats payments as amounts ONLY (no dates) for the REMEDIAL_PAYMENT_BALANCES file.
        Dates are stored separately in PAYMENT_DATES_LOG.xlsx.
        
        Args:
            payments: List of PaymentRecord objects
            
        Returns:
            List of amount values only (one per payment)
            [amount1, amount2, ...]
        """
        formatted = []
        
        for payment in payments:
            # Add amount as float for Excel - NO DATE in remedial balances file
            formatted.append(float(payment.amount))
        
        return formatted
    
    def update_row(self, row_index: int, data: dict, worksheet) -> dict:
        """
        Update a specific learner row in the worksheet.
        
        Args:
            row_index: The row number to update (1-indexed)
            data: Dictionary with keys 'name', 'admission_number', 'total_owed', 'payments'
            worksheet: The openpyxl worksheet object
            
        Returns:
            dict with 'success' (bool) and 'message' (str) keys
            
        Note:
            This method updates the learner's basic information and completely
            replaces their payment history. Old payment data is cleared before
            writing new data to prevent data corruption.
            Handles merged cells by checking before writing.
        """
        try:
            logger.debug(f"Updating row {row_index} with data for {data['admission_number']}")
            
            # Determine column structure based on sheet name
            is_form4_2026 = "FORM 4 2026" in worksheet.title.upper() or "FORM 4A" in worksheet.title.upper()
            
            if is_form4_2026:
                # Form 4 2026 structure:
                # Column D (4): Admission Number
                # Column E (5): Name
                # Column G (7): Total Owed
                # Column R (18): JAN, S(19): FEB, T(20): MAR, U(21): APR, V(22): MAY, W(23): JUN, X(24): JUL
                adm_col = 4
                name_col = 5
                total_col = 7
                payment_start_col = 18  # Column R = January
            else:
                # Standard structure:
                # Column B (2): Admission Number
                # Column C (3): Name
                # Column E (5): Total Owed
                # Column H (8): JAN, I(9): FEB, J(10): MAR, K(11): APR, L(12): MAY, M(13): JUN,
                # N(14): JUL, O(15): AUG, P(16): SEP, Q(17): OCT, R(18): NOV
                adm_col = 2
                name_col = 3
                total_col = 5
                payment_start_col = 8  # Column H = January
            
            # Helper function to check if a cell position is in a merged range
            def is_merged_cell(row, col):
                from openpyxl.utils import get_column_letter
                cell_coord = f"{get_column_letter(col)}{row}"
                for merged_range in worksheet.merged_cells.ranges:
                    if cell_coord in merged_range:
                        return True
                return False
            
            # Helper function to safely write to a cell (skip if merged)
            def safe_write_cell(row, col, value):
                try:
                    # Check if position is merged BEFORE accessing the cell
                    if is_merged_cell(row, col):
                        logger.warning(f"Skipping merged cell at row {row}, col {col}")
                        return False
                    
                    # Now safe to access and write
                    cell = worksheet.cell(row=row, column=col)
                    cell.value = value
                    return True
                except Exception as e:
                    logger.warning(f"Could not write to cell at row {row}, col {col}: {e}")
                    return False
            
            # Update basic learner information in correct columns
            safe_write_cell(row_index, adm_col, data['admission_number'])
            safe_write_cell(row_index, name_col, data['name'])
            safe_write_cell(row_index, total_col, data['total_owed'])
            
            # Clear existing payment month columns (Jan=payment_start_col through Nov/Jul)
            # Standard sheet: H(8) through R(18) = 11 months
            # Form 4 2026: R(18) through X(24) = 7 months
            if is_form4_2026:
                clear_end_col = payment_start_col + 7  # JAN through JUL (7 months)
            else:
                clear_end_col = payment_start_col + 11  # JAN through NOV (11 months)
            
            for col in range(payment_start_col, clear_end_col):
                try:
                    if is_merged_cell(row_index, col):
                        continue
                    cell = worksheet.cell(row=row_index, column=col)
                    if cell.value is not None:
                        cell.value = None
                except Exception:
                    pass
            
            # Month-to-column mapping: month number -> column offset from payment_start_col
            # January=0, February=1, March=2, ... November=10
            MONTH_OFFSET = {1: 0, 2: 1, 3: 2, 4: 3, 5: 4, 6: 5, 7: 6, 8: 7, 9: 8, 10: 9, 11: 10}
            
            # Form 4 2026 only goes Jan-Jul (7 months): 1->0, 2->1, 3->2, 4->3, 5->4, 6->5, 7->6
            FORM4_MONTH_OFFSET = {1: 0, 2: 1, 3: 2, 4: 3, 5: 4, 6: 5, 7: 6}
            
            month_offset_map = FORM4_MONTH_OFFSET if is_form4_2026 else MONTH_OFFSET
            
            # Write payment records to the correct month columns
            payment_records = data.get('payment_records', [])
            if payment_records:
                # Use full payment records with timestamps for month-based placement
                for payment in payment_records:
                    try:
                        ts = payment.timestamp
                        if ts and hasattr(ts, 'month'):
                            month = ts.month
                        else:
                            # Try to extract month from string timestamp
                            from datetime import datetime as dt
                            if isinstance(ts, str):
                                try:
                                    ts = dt.fromisoformat(ts)
                                    month = ts.month
                                except Exception:
                                    month = None
                            else:
                                month = None
                        
                        if month and month in month_offset_map:
                            col = payment_start_col + month_offset_map[month]
                            # Add to existing value in that month column (multiple payments in same month)
                            try:
                                if not is_merged_cell(row_index, col):
                                    cell = worksheet.cell(row=row_index, column=col)
                                    existing = cell.value
                                    if existing and isinstance(existing, (int, float)):
                                        cell.value = float(existing) + float(payment.amount)
                                    else:
                                        cell.value = float(payment.amount)
                            except Exception as e:
                                logger.warning(f"Could not write payment to col {col}: {e}")
                        else:
                            # Unknown month or out of range: fall back to sequential from start
                            logger.warning(f"Payment month {month} not mappable for {data['admission_number']}, using sequential")
                            payments_flat = data.get('payments', [])
                            col_index = payment_start_col
                            for value in payments_flat:
                                safe_write_cell(row_index, col_index, value)
                                col_index += 1
                            break  # Don't repeat; flat write handles all
                    except Exception as e:
                        logger.warning(f"Error placing payment by month: {e}")
            else:
                # Fallback: sequential write (backward compatibility)
                payments_flat = data.get('payments', [])
                col_index = payment_start_col
                for value in payments_flat:
                    safe_write_cell(row_index, col_index, value)
                    col_index += 1
            
            logger.debug(f"Successfully updated row {row_index} with {len(payment_records)} payment records")
            
            return {
                'success': True,
                'message': f'Successfully updated learner {data["admission_number"]}'
            }
            
        except Exception as e:
            error_msg = f'Failed to update row {row_index}: {str(e)}'
            logger.error(error_msg)
            return {
                'success': False,
                'message': error_msg
            }
    
    def _find_learner_row(self, worksheet, admission_number: str, sheet_name: str = "", stream_suffix: str = None) -> int:
        """
        Find the row index for a learner by admission number.
        
        Searches through the worksheet starting from row 5 (where data begins)
        and looks for a matching admission number in the correct column and row range.
        
        Args:
            worksheet: The openpyxl worksheet object
            admission_number: The admission number to search for
            sheet_name: Name of the sheet (for determining row ranges)
            stream_suffix: Stream suffix like "- Achievers" or "- Champions" to narrow search
            
        Returns:
            Row index (1-indexed) if found, None otherwise
            
        Note:
            The search starts at row 5 because rows 1-4 typically contain
            headers and metadata in the spreadsheet format.
            If stream_suffix is provided, only searches within that stream's row range.
        """
        # Determine which column contains admission numbers based on sheet
        is_form4_2026 = "FORM 4 2026" in sheet_name.upper() or "FORM 4A" in sheet_name.upper()
        adm_col_index = 3 if is_form4_2026 else 1  # Column D (index 3) for Form 4, Column B (index 1) for others
        
        # Define row ranges for each sheet and stream (same as in excel_reader.py)
        stream_ranges = {
            'GRADE 10 2026': {
                '- Achievers': (6, 58),
                '- Champions': (73, 100),
            },
            'FORM 2 2025': {
                '- Achievers': (5, 75),
                '- Champions': (104, 171),
            },
            'FORM 3 2026': {
                '- Achievers': (5, 75),
                '- Champions': (104, 171),
            },
            'FORM 3 2025': {
                '- Achievers': (5, 66),
                '- Champions': (75, 124),
            },
            'FORM 4 2026': {
                '- Achievers': (5, 68),
                '- Champions': (76, 125),
            },
        }
        
        # Determine search range
        min_row = 5
        max_row = worksheet.max_row
        
        if stream_suffix and sheet_name in stream_ranges:
            if stream_suffix in stream_ranges[sheet_name]:
                min_row, max_row = stream_ranges[sheet_name][stream_suffix]
                logger.debug(f"Searching for {admission_number} in {sheet_name} {stream_suffix} rows {min_row}-{max_row}")
        
        # Iterate through rows in the determined range
        for row_idx, row in enumerate(worksheet.iter_rows(min_row=min_row, max_row=max_row, values_only=False), start=min_row):
            if len(row) > adm_col_index:
                cell_value = row[adm_col_index].value  # Admission number column
                # Compare as strings after stripping whitespace to handle formatting differences
                if cell_value and str(cell_value).strip() == str(admission_number).strip():
                    logger.debug(f"Found learner {admission_number} at row {row_idx}")
                    return row_idx
        
        # Learner not found in the worksheet
        logger.warning(f"Learner {admission_number} not found in {sheet_name} {stream_suffix or ''} (searched rows {min_row}-{max_row})")
        return None

    def _log_payment_dates(self, learner: Learner, log_file: str) -> None:
        """
        Log payment dates to PAYMENT_DATES_LOG.xlsx.
        
        Writes each payment record with its date to the log file.
        Columns: Date, Admission Number, Learner Name, Amount, Month, Sheet/Class, Timestamp
        
        Args:
            learner: The Learner object with payment records
            log_file: Path to PAYMENT_DATES_LOG.xlsx
        """
        try:
            import os
            from openpyxl import load_workbook, Workbook
            
            # Load or create the log file
            if os.path.exists(log_file):
                wb = load_workbook(log_file)
                if 'Payment Log' in wb.sheetnames:
                    ws = wb['Payment Log']
                else:
                    ws = wb.active
            else:
                wb = Workbook()
                ws = wb.active
                ws.title = 'Payment Log'
                # Write headers
                ws.append(['Date', 'Admission Number', 'Learner Name', 'Amount', 'Month', 'Sheet/Class', 'Timestamp'])
            
            # Get existing entries to avoid duplicates
            # Build a set of (admission_number, amount, timestamp_str) already logged
            existing = set()
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[1] and row[3] and row[6]:
                    existing.add((str(row[1]), str(row[3]), str(row[6])))
            
            # Append new payment records
            for payment in learner.payment_records:
                ts = payment.timestamp
                ts_str = ts.strftime('%Y-%m-%d %H:%M:%S') if ts else ''
                amount_str = str(float(payment.amount))
                key = (str(learner.admission_number), amount_str, ts_str)
                
                if key not in existing:
                    date_str = ts.strftime('%Y-%m-%d') if ts else ''
                    month_str = ts.strftime('%b').upper() if ts else ''
                    ws.append([
                        date_str,
                        learner.admission_number,
                        learner.name,
                        float(payment.amount),
                        month_str,
                        learner.source_sheet or '',
                        ts_str
                    ])
                    existing.add(key)
            
            wb.save(log_file)
            wb.close()
            logger.info(f"Payment dates logged to {log_file}")
            
        except Exception as e:
            # Don't fail the main operation if logging fails
            logger.warning(f"Could not log payment dates: {str(e)}")
